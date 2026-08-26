#!/usr/bin/env python3
"""Render every worksheet to a lightweight PNG for visual QA without Excel automation."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
BOOK = ROOT / "outputs" / "mango_bd_pilot_20260824" / "Mango_Labs_AI_BD_Pilot.xlsx"
OUT = ROOT / "outputs" / "mango_bd_pilot_20260824" / "previews"
MAX_WIDTH = 5600
MAX_HEIGHT = 7600


def font(size: int, bold: bool = False):
    candidates = [
        "/System/Library/Fonts/Supplemental/Arial Bold.ttf" if bold else "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/System/Library/Fonts/Supplemental/Helvetica.ttc",
    ]
    for candidate in candidates:
        if Path(candidate).exists():
            return ImageFont.truetype(candidate, size)
    return ImageFont.load_default()


def color(cell, default: str = "FFFFFF") -> str:
    fill = cell.fill
    if fill and fill.fill_type == "solid" and fill.fgColor.type == "rgb":
        raw = fill.fgColor.rgb or default
        return raw[-6:]
    return default


def text_color(cell, default: str = "1F2937") -> str:
    if cell.font.color and cell.font.color.type == "rgb" and cell.font.color.rgb:
        return cell.font.color.rgb[-6:]
    return default


def shorten(value, chars: int) -> str:
    if value is None:
        return ""
    text = str(value).replace("\n", " ")
    return text if len(text) <= chars else text[: max(1, chars - 1)] + "…"


def render_sheet(ws, path: Path) -> None:
    max_row = max(1, ws.max_row)
    max_col = max(1, ws.max_column)
    col_widths = []
    for col in range(1, max_col + 1):
        letter = ws.cell(4, col).column_letter
        width = ws.column_dimensions[letter].width or 12
        col_widths.append(int(min(max(width * 6.0, 56), 310)))
    row_heights = []
    for row in range(1, max_row + 1):
        height = ws.row_dimensions[row].height or 18
        row_heights.append(int(min(max(height * 1.15, 24), 130)))

    total_w = sum(col_widths) + 2
    total_h = sum(row_heights) + 2
    scale = min(1.0, MAX_WIDTH / total_w, MAX_HEIGHT / total_h)
    canvas = Image.new("RGB", (total_w, total_h), "#FFFFFF")
    draw = ImageDraw.Draw(canvas)
    y = 1
    for row in range(1, max_row + 1):
        x = 1
        rh = row_heights[row - 1]
        for col in range(1, max_col + 1):
            cw = col_widths[col - 1]
            cell = ws.cell(row, col)
            draw.rectangle((x, y, x + cw, y + rh), fill="#" + color(cell), outline="#D8DEE6", width=1)
            size = 20 if row <= 2 else (12 if row in (3, 5) else 11)
            cell_font = font(size, bold=bool(cell.font.bold or row in (1, 5)))
            char_limit = max(7, int(cw / (size * 0.55)))
            raw = cell.value
            value = shorten(raw, char_limit * max(1, int(rh / (size + 3))))
            draw.multiline_text((x + 5, y + 4), value, fill="#" + text_color(cell), font=cell_font, spacing=2)
            x += cw
        y += rh
    if scale < 1.0:
        canvas = canvas.resize((int(total_w * scale), int(total_h * scale)), Image.Resampling.LANCZOS)
    canvas.save(path)


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    workbook = load_workbook(BOOK, data_only=False)
    for index, ws in enumerate(workbook.worksheets, 1):
        safe = ws.title.lower().replace(" ", "_")
        path = OUT / f"{index:02d}_{safe}.png"
        render_sheet(ws, path)
        print(path)


if __name__ == "__main__":
    main()
