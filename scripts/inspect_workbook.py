#!/usr/bin/env python3
"""Structural and formula QA for the generated workbook."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
BOOK = ROOT / "outputs" / "mango_bd_pilot_20260824" / "Mango_Labs_AI_BD_Pilot.xlsx"
EXPECTED = {
    "Dashboard": 1,
    "Scoring": 5,
    "Projects": 12,
    "Priority Actions": 8,
    "Creators": 13,
    "Relationships": 68,
    "Entities": 73,
    "Evidence": 53,
    "GTM Cases": 6,
    "README": 8,
}


def main() -> None:
    wb = load_workbook(BOOK, data_only=False)
    assert wb.sheetnames == list(EXPECTED), wb.sheetnames
    formulas = []
    errors = []
    comments = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if isinstance(value, str) and value.startswith("="):
                    formulas.append(f"{ws.title}!{cell.coordinate}")
                if isinstance(value, str) and value.strip().upper() in {"#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!"}:
                    errors.append(f"{ws.title}!{cell.coordinate}={value}")
                if cell.comment:
                    comments.append((f"{ws.title}!{cell.coordinate}", cell.comment.author))
    assert not errors, errors
    assert all(author == "User" for _, author in comments), comments
    assert len(formulas) >= 36, len(formulas)
    assert len(wb["Projects"].tables) == 1
    assert len(wb["Dashboard"]._charts) == 1
    assert len(wb["Priority Actions"].data_validations.dataValidation) == 1
    print({"sheets": len(wb.sheetnames), "formulas": len(formulas), "comments": len(comments), "formula_errors": len(errors)})
    for name in wb.sheetnames:
        ws = wb[name]
        print(f"{name}: {ws.max_row} rows x {ws.max_column} cols")


if __name__ == "__main__":
    main()
