#!/usr/bin/env python3
"""Build the two-sheet Head AI KOL and supplier shortlist workbook."""

from __future__ import annotations

import csv
import json
import re
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
KOL_CSV = ROOT / "head_ai_kol_shortlist.csv"
HTML_DATA = ROOT / "head_ai_kol_and_supplier_shortlist.html"
OUTPUT_DIR = ROOT / "outputs" / "ai_kol_new_outreach_20260906_expanded"
OUTPUT_FILE = OUTPUT_DIR / "head_ai_kol_and_supplier_new_outreach_shortlist.xlsx"

NAVY = "172033"
TEAL = "167D7F"
CORAL = "F15A3A"
WHITE = "FFFFFF"
DARK = "202733"
MUTED = "667085"
LINK = "0563C1"
PALE_GRAY = "F5F7FA"
PALE_TEAL = "E7F5F3"
PALE_ORANGE = "FFF1E8"
PALE_YELLOW = "FFF7D6"
PALE_RED = "FDECEC"
MID_GRAY = "D7DEE8"
THIN = Side(style="thin", color=MID_GRAY)


EXTERNAL_KOLS = [
    {
        "平台": "YouTube",
        "账号": "Matt Wolfe（@mreflow）",
        "主页链接": "https://www.youtube.com/@mreflow",
        "粉丝数": "1.00M",
        "报价": "待询价",
        "联系方式": "mattwolfe@smoothmedia.co；https://tally.so/r/nrBVlp（品牌赞助表单）；Smooth Media",
        "跟进时间": "未跟进",
        "备注": (
            "S；本地 creator ID 330；本地 Pika 记录仅为 mention，未计现金赞助。"
            "外部 SponsorRadar 聚合 22 品牌/31 笔赞助记录，最新至 2026-08-28；官方频道公开 Smooth Media 商务邮箱与表单。"
            "Mango 无报价/回复；建议经 Smooth Media 询价 AI 工具首发与目录型曝光。"
            "外部补充核验 2026-09-06：https://sponsorradar.com/channels/mreflow；https://mattwolfe.com/"
        ),
    },
    {
        "平台": "YouTube",
        "账号": "Tech With Tim（@techwithtim）",
        "主页链接": "https://www.youtube.com/@techwithtim",
        "粉丝数": "2.08M",
        "报价": "待询价",
        "联系方式": "contact@kernelmanagement.com；tim@techwithtim.net；Kernel Management / Reign Maker Talent",
        "跟进时间": "未跟进",
        "备注": (
            "S；SponsorRadar 聚合 37 品牌/87 笔，含 Hostinger、DataCamp、Boot.dev 等开发者品牌；"
            "公开资料显示由 Kernel Management / Reign Maker Talent 承接科技创作者商务。适合 AI agent、coding 与 developer tools。"
            "Mango 无报价/回复；建议先向 Kernel 询单人价及 developer roster。"
            "外部补充核验 2026-09-06：https://sponsorradar.com/channels/techwithtim；https://kernelmanagement.com/"
        ),
    },
    {
        "平台": "YouTube",
        "账号": "Ishan Sharma（@ishansharma7390）",
        "主页链接": "https://www.youtube.com/@ishansharma7390",
        "粉丝数": "2.18M",
        "报价": "待询价",
        "联系方式": "collab@helloishan.com；MarkitUp Media",
        "跟进时间": "未跟进",
        "备注": (
            "S；SponsorRadar 聚合 21 品牌/29 笔，Anthropic 3、Coursera 3、Replit 2；官方站列出 Gamma、Gemini、Higgsfield、Lindy 等合作。"
            "印度英语市场头部，适合职业/创业/AI 工具；MarkitUp 为其自有内容 agency，但代理其他 creator 的权利待核验。"
            "Mango 无报价/回复。外部补充核验 2026-09-06：https://helloishan.com/；https://sponsorradar.com/channels/ishansharma7390"
        ),
    },
    {
        "平台": "YouTube",
        "账号": "Fireship（@Fireship）",
        "主页链接": "https://www.youtube.com/@Fireship",
        "粉丝数": "4.26M",
        "报价": "待询价",
        "联系方式": "https://www.electrify.video/our-brand-partnerships（Electrify）",
        "跟进时间": "未跟进",
        "备注": (
            "S；SponsorRadar 聚合 43 品牌/97 笔，含 Brilliant、CodeRabbit、Hostinger、Railway；开发者信任度和品牌认知价值高。"
            "Electrify 官方将 Fireship 列为旗下 creator-founded brand，并提供多创作者组合采购；适合品牌认知，长教程转化需先确认。"
            "Mango 无报价/回复；建议从 Electrify 机构入口询价。外部补充核验 2026-09-06："
            "https://sponsorradar.com/channels/fireship；https://www.electrify.video/our-brands"
        ),
    },
    {
        "平台": "YouTube",
        "账号": "Futurepedia（@futurepedia_io）",
        "主页链接": "https://www.youtube.com/@futurepedia_io",
        "粉丝数": "748K",
        "报价": "待询价",
        "联系方式": "https://www.futurepedia.io/youtube-interest（官方广告表单）",
        "跟进时间": "未跟进",
        "备注": (
            "S；AI 工具发行型头部；官方广告页覆盖 Futurepedia、Skill Leap AI、Howfinity，可谈植入、专场、课程和 webinar，且明确不接 affiliate offer。"
            "SponsorRadar 聚合 7 个赞助品牌。Mango 无报价/回复；建议以三频道组合和单频道净价同时询价。"
            "外部补充核验 2026-09-06：https://www.futurepedia.io/youtube-interest；https://sponsorradar.com/channels/futurepedia-io"
        ),
    },
    {
        "平台": "YouTube",
        "账号": "Lenny’s Podcast（@lennyspodcast）",
        "主页链接": "https://www.youtube.com/@lennyspodcast",
        "粉丝数": "633K",
        "报价": "待询价",
        "联系方式": "podcast@lennyrachitsky.com",
        "跟进时间": "未跟进",
        "备注": (
            "S；SponsorRadar 聚合 17 品牌/59 笔；公开节目说明提供 sponsor 邮箱，受众集中产品经理、创始人与 B2B 买方。"
            "适合 Cursor、Gamma、Replit、Lovable、n8n 等高客单/决策者触达，不以底部漏斗长演示为主。"
            "sponsorship.so 的 47 品牌/261 次为提及密度，不等同全部现金赞助。Mango 无报价/回复。"
            "外部补充核验 2026-09-06：https://sponsorradar.com/channels/lennyspodcast"
        ),
    },
    {
        "平台": "YouTube",
        "账号": "Theo — t3.gg（@t3dotgg）",
        "主页链接": "https://www.youtube.com/@t3dotgg",
        "粉丝数": "563K",
        "报价": "待询价",
        "联系方式": "youtube@t3.gg；https://t3.gg/sponsor-me",
        "跟进时间": "未跟进",
        "备注": (
            "S；官方赞助页称月均 3M views、已帮助 40+ 品牌；SponsorRadar 聚合 61 品牌/355 笔。"
            "开发者意见领袖，适合 coding agent、AI infra 与 developer tools；不适合泛生产力。"
            "Mango 无报价/回复；建议直接走官方 sponsor 入口。外部补充核验 2026-09-06："
            "https://t3.gg/sponsor-me；https://sponsorradar.com/channels/t3dotgg"
        ),
    },
    {
        "平台": "YouTube",
        "账号": "AI Search（@theaisearch）",
        "主页链接": "https://www.youtube.com/@theaisearch",
        "粉丝数": "726K",
        "报价": "待询价",
        "联系方式": "https://www.youtube.com/@theaisearch/about（公开 business inquiry）",
        "跟进时间": "未跟进",
        "备注": (
            "S；SponsorRadar 聚合 53 品牌/192 笔，Abacus AI、HubSpot、Higgsfield 等重复出现；AI 工具评测商业密度高。"
            "适合多工具测评和效果向投放，但需抽查近 3 条，区分现金赞助、联盟链接和自然提及。Mango 无报价/回复。"
            "外部补充核验 2026-09-06：https://sponsorradar.com/channels/theaisearch；"
            "https://sponsorship.so/top-influencers/ai（后者为品牌提及统计）"
        ),
    },
    {
        "平台": "YouTube",
        "账号": "MattVidPro（@MattVidPro）",
        "主页链接": "https://www.youtube.com/@MattVidPro",
        "粉丝数": "304K",
        "报价": "待询价",
        "联系方式": "mattvidpro@smoothmedia.co；https://tally.so/r/3xdz4E；Smooth Media",
        "跟进时间": "未跟进",
        "备注": (
            "S；本地数据库有 Verda 明确付费赞助证据（2026-04-03，observed_unreviewed）；SponsorRadar 聚合 38 品牌/62 笔，含 Invideo、Recraft、Verda。"
            "AI 视频/图像生成垂直，且由 Smooth Media 提供商务入口；适合与 Matt Wolfe 打包询价。Mango 无报价/回复。"
            "外部补充核验 2026-09-06：https://sponsorradar.com/channels/mattvidpro"
        ),
    },
    {
        "平台": "YouTube",
        "账号": "Tina Huang（@TinaHuang1）",
        "主页链接": "https://www.youtube.com/@TinaHuang1",
        "粉丝数": "1.30M",
        "报价": "待询价",
        "联系方式": "hellotinah@gmail.com",
        "跟进时间": "未跟进",
        "备注": (
            "S；SponsorRadar 聚合 19 品牌/72 笔，含 HubSpot、Brilliant，以及 Perplexity、Genspark、Bolt、Wispr 等 AI 工具。"
            "数据科学/AI/职业学习受众，适合职业向 AI campaign；硬核模型评测适配较弱。Mango 无报价/回复。"
            "外部补充核验 2026-09-06：https://sponsorradar.com/channels/tinahuang1；公开商务邮箱来自频道公开资料。"
        ),
    },
    {
        "平台": "YouTube",
        "账号": "Ali Abdaal（@aliabdaal）",
        "主页链接": "https://www.youtube.com/@aliabdaal",
        "粉丝数": "6.68M",
        "报价": "待询价",
        "联系方式": "https://aliabdaal.com/contact/；https://www.passionfroot.me/ali-abdaal",
        "跟进时间": "未跟进",
        "备注": (
            "A；生产力超头部而非纯 AI 垂直；SponsorRadar 聚合 22 品牌/41 笔，近年可见 Perplexity、Trading 212、Epidemic 等。"
            "适合知识工作者/创业品牌背书，预计门槛高，需大预算和明确 brief。Mango 无报价/回复；建议官方表单或 Passionfroot 首次询价。"
            "外部补充核验 2026-09-06：https://sponsorradar.com/channels/aliabdaal；https://aliabdaal.com/contact/"
        ),
    },
    {
        "平台": "YouTube",
        "账号": "Skill Leap AI（@SkillLeapAI）",
        "主页链接": "https://www.youtube.com/@SkillLeapAI",
        "粉丝数": "338K",
        "报价": "待询价",
        "联系方式": "saj@skillleap.ai；https://www.futurepedia.io/youtube-interest",
        "跟进时间": "未跟进",
        "备注": (
            "A；Futurepedia 同体系的 AI 工具教程频道；官方广告页支持植入、专场和 custom course，并明确不接 affiliate offer。"
            "适合教程转化和与 Futurepedia 组合采购。Mango 无报价/回复；建议通过统一表单询价并确认单频道/组合价。"
            "外部补充核验 2026-09-06：https://www.futurepedia.io/youtube-interest"
        ),
    },
    {
        "平台": "YouTube",
        "账号": "The AI Advantage（@aiadvantage）",
        "主页链接": "https://www.youtube.com/@aiadvantage",
        "粉丝数": "478K",
        "报价": "待询价",
        "联系方式": "https://www.youtube.com/@aiadvantage/about；https://myaiadvantage.com/",
        "跟进时间": "未跟进",
        "备注": (
            "A；sponsorship.so 记录 32 品牌/111 次提及，但该口径不是全部现金赞助；公开视频说明可确认 LTX Studio 赞助。"
            "未找到足够证据支持“Luma AI 明确现金赞助”，已纠正为待核验。适合实用 AI 教程/周更新闻，询价前抽查近 3 条。"
            "Mango 无报价/回复。外部补充核验 2026-09-06：https://sponsorship.so/top-influencers/ai"
        ),
    },
    {
        "平台": "YouTube",
        "账号": "Aurelius Tjin（@AureliusTjin）",
        "主页链接": "https://www.youtube.com/@AureliusTjin",
        "粉丝数": "562K",
        "报价": "待询价",
        "联系方式": "partnerships@aureliustjin.com",
        "跟进时间": "未跟进",
        "备注": (
            "A；SponsorRadar 聚合 19 品牌/45 笔，含 Hostinger、Canva、Gamma；创作者工具/设计 AI 方向匹配，商务邮箱完整。"
            "Mango 无报价/回复；适合作为垂直价格带样本和教程型转化候选。"
            "外部补充核验 2026-09-06：https://sponsorradar.com/channels/aureliustjin；https://www.aureliustjin.com/"
        ),
    },
    {
        "平台": "YouTube",
        "账号": "Traversy Media（@TraversyMedia）",
        "主页链接": "https://www.youtube.com/@TraversyMedia",
        "粉丝数": "2.42M",
        "报价": "待询价",
        "联系方式": "https://www.youtube.com/@TraversyMedia/about；https://www.traversymedia.com/",
        "跟进时间": "未跟进",
        "备注": (
            "A；本地 Mango 建联总表存在 X 账号记录且无报价/回复标记；SponsorRadar 聚合 16 品牌/26 笔，含 CodeRabbit、Coding with AI。"
            "老牌开发者教育频道，适合 coding/AI developer tools；AI 垂直度低于 Fireship、Theo。Mango 无报价/回复。"
            "外部补充核验 2026-09-06：https://sponsorradar.com/channels/traversymedia"
        ),
    },
]


def load_kols() -> list[dict[str, str]]:
    with KOL_CSV.open("r", encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    expected = ["平台", "账号", "主页链接", "粉丝数", "报价", "联系方式", "跟进时间", "备注"]
    if not rows or list(rows[0]) != expected:
        raise ValueError(f"Unexpected KOL columns: {list(rows[0]) if rows else []}")
    return rows


def load_embedded_data() -> dict:
    text = HTML_DATA.read_text(encoding="utf-8")
    prefix = "const DATA = "
    start = text.index(prefix) + len(prefix)
    end = text.index(";\n    const state", start)
    return json.loads(text[start:end])


def follower_number(value: str) -> float:
    text = str(value or "").strip().upper().replace(",", "")
    if text.endswith("M"):
        return float(text[:-1]) * 1_000_000
    if text.endswith("K"):
        return float(text[:-1]) * 1_000
    try:
        return float(text)
    except ValueError:
        return 0


def filter_new_kols(rows: list[dict[str, str]], data: dict) -> tuple[list[dict[str, str]], list[dict[str, str]]]:
    status_by_url = {row["profileUrl"]: row["relationshipStatus"] for row in data["kolRows"]}
    kept: list[dict[str, str]] = []
    removed: list[dict[str, str]] = []
    for row in rows:
        relationship = status_by_url.get(row["主页链接"], "待核验")
        has_quote = not row["报价"].startswith("待询价")
        affiliate_only = "affiliate-only" in row["备注"].lower()
        if relationship != "待建联" or has_quote or affiliate_only:
            removed.append(row)
        else:
            kept.append(row)
    return kept, removed


def merge_external_kols(rows: list[dict[str, str]]) -> list[dict[str, str]]:
    """Prepend user-requested, externally verified leads and deduplicate by profile URL."""
    combined = [*EXTERNAL_KOLS, *rows]
    seen: set[str] = set()
    deduped: list[dict[str, str]] = []
    for row in combined:
        profile = row["主页链接"].rstrip("/").lower()
        if profile in seen:
            continue
        seen.add(profile)
        deduped.append(row)
    return deduped


def kol_child(row: dict[str, str], entity_id: str, relationship: str) -> dict[str, object]:
    handle_match = re.search(r"@([A-Za-z0-9_.-]+)", f"{row['账号']} {row['主页链接']}")
    return {
        "creatorId": entity_id,
        "creator": re.sub(r"（@[^）]+）", "", row["账号"]).strip(),
        "platform": row["平台"],
        "handle": handle_match.group(1) if handle_match else row["账号"],
        "profileUrl": row["主页链接"],
        "followers": row["粉丝数"],
        "contact": row["联系方式"],
        "relationship": relationship,
    }


def build_new_supplier_rows(data: dict, kols: list[dict[str, str]]) -> list[dict]:
    original = {supplier["name"]: supplier for supplier in data["supplierRows"]}

    claryo = dict(original["Claryo Media"])
    claryo.update(
        priority="A",
        relationshipStatus="未建联",
        nextAction="先核验 creator 授权和交付案例，再索取可供 Mango 采购的 AI roster",
    )

    passionfroot = dict(original["Passionfroot"])
    booking_candidates = [
        child for child in passionfroot["linkedKols"]
        if child["creator"] != "Linus Tech Tips" and follower_number(child["followers"]) >= 10_000
    ][:8]
    passionfroot.update(
        priority="A",
        coverage=len(booking_candidates),
        linkedKols=booking_candidates,
        relationshipStatus="未建联（公开预订入口）",
        nextAction="按具体 campaign brief 从候选中选择 5-8 个发起 request，不按完整平台库存报覆盖",
    )

    lmg = dict(original["Linus Media Group"])
    lmg.update(
        priority="S",
        relationshipStatus="未建联（公开合作入口）",
        nextAction="仅在大型 AI infra、hardware 或 enterprise 预算明确时首次询价",
    )

    by_handle = {}
    for row in kols:
        match = re.search(r"@([A-Za-z0-9_.-]+)", f"{row['账号']} {row['主页链接']}")
        if match:
            by_handle[match.group(1).lower()] = row

    sonny = by_handle["sonnysangha"]
    blog_with_ben = by_handle["blogwithben"]
    tiago = by_handle["tiagoforte"]
    matt_wolfe = by_handle["mreflow"]
    mattvidpro = by_handle["mattvidpro"]
    tech_with_tim = by_handle["techwithtim"]
    ishan = by_handle["ishansharma7390"]
    fireship = by_handle["fireship"]
    futurepedia = by_handle["futurepedia_io"]
    skill_leap = by_handle["skillleapai"]
    ali_abdaal = by_handle["aliabdaal"]

    if not any(child["creator"] == "Ali Abdaal" for child in booking_candidates):
        booking_candidates.append(
            kol_child(ali_abdaal, "external:aliabdaal", "公开 storefront，可预订但非独家代理证明")
        )
        passionfroot.update(coverage=len(booking_candidates), linkedKols=booking_candidates)

    smooth_media = {
        "id": "smooth-media",
        "priority": "S",
        "name": "Smooth Media",
        "type": "Knowledge-creator management and brand partnerships；官方称管理 50+ creators",
        "contact": "hi@smoothmedia.co；mattwolfe@smoothmedia.co；mattvidpro@smoothmedia.co",
        "officialUrl": "https://smoothmedia.co/about",
        "relationshipStatus": "未建联",
        "verification": "高",
        "nextAction": "以 Matt Wolfe + MattVidPro 双账号询价切入，索取 AI/tech roster、组合净价、agency/sub-affiliate/whitelisting 权限",
        "coverage": 2,
        "linkedKols": [
            kol_child(matt_wolfe, "creator:330", "官方频道明确给出 Smooth Media 商务邮箱"),
            kol_child(mattvidpro, "external:mattvidpro", "官方频道明确给出 Smooth Media 商务邮箱与品牌表单"),
        ],
    }

    kernel = {
        "id": "kernel-reign-maker",
        "priority": "S",
        "name": "Kernel Management / Reign Maker Talent",
        "type": "Tech/coding creator talent management；公开 roster 总量未披露",
        "contact": "contact@kernelmanagement.com；tim@techwithtim.net",
        "officialUrl": "https://kernelmanagement.com/",
        "relationshipStatus": "未建联",
        "verification": "高",
        "nextAction": "先询 Tech With Tim 单人价，再索取可批量采购的 developer/AI creator roster 与授权范围",
        "coverage": 1,
        "linkedKols": [
            kol_child(tech_with_tim, "external:techwithtim", "官方/公开资料确认 Kernel 与 Reign Maker 的人才管理路径"),
        ],
    }

    markitup = {
        "id": "markitup-media",
        "priority": "A",
        "name": "MarkitUp Media",
        "type": "Creator-owned content marketing and YouTube management agency；对外 creator 代理权待核验",
        "contact": "collab@helloishan.com；tanisha@helloishan.com",
        "officialUrl": "https://markitup.in/",
        "relationshipStatus": "未建联",
        "verification": "中高",
        "nextAction": "先以 Ishan Sharma 品牌合作询价，确认是否可提供其他南亚 AI/tech creators、批量采购与本地执行",
        "coverage": 1,
        "linkedKols": [
            kol_child(ishan, "external:ishansharma7390", "Ishan 官方网站列其为共同创办的内容 agency；他人代理权限待核验"),
        ],
    }

    electrify = {
        "id": "electrify",
        "priority": "S",
        "name": "Electrify",
        "type": "Creator-founded media group；官网列 9 个 creator brands，可执行 multi-creator campaigns",
        "contact": "https://www.electrify.video/our-brand-partnerships",
        "officialUrl": "https://www.electrify.video/our-brands",
        "relationshipStatus": "未建联",
        "verification": "高",
        "nextAction": "先询 Fireship 档期与产品适配，再索取 9 个官方品牌中的 tech/AI 可采库存、组合价和授权；注意其不是开放 MCN",
        "coverage": 1,
        "linkedKols": [
            kol_child(fireship, "external:fireship", "Electrify 官方品牌页列名，并明确承接赞助与多创作者组合"),
        ],
    }

    futurepedia_network = {
        "id": "futurepedia-network",
        "priority": "S",
        "name": "Futurepedia / Skill Leap AI",
        "type": "AI education media network；官方广告页覆盖 3 个频道",
        "contact": "https://www.futurepedia.io/youtube-interest；saj@skillleap.ai",
        "officialUrl": "https://www.futurepedia.io/youtube-interest",
        "relationshipStatus": "未建联",
        "verification": "高",
        "nextAction": "同时询 Futurepedia、Skill Leap AI、Howfinity 的单频道/组合净价、交付形式和排他期；官方明确不接 affiliate offer",
        "coverage": 3,
        "linkedKols": [
            kol_child(futurepedia, "external:futurepedia", "官方广告页列入同一可采购体系"),
            kol_child(skill_leap, "external:skillleapai", "官方广告页列入同一可采购体系"),
            {
                "creatorId": "external:howfinity",
                "creator": "Howfinity",
                "platform": "YouTube",
                "handle": "Howfinity",
                "profileUrl": "https://www.youtube.com/@Howfinity",
                "followers": "1.10M",
                "contact": "https://www.futurepedia.io/youtube-interest",
                "relationship": "官方广告页列入三频道组合；未单列入 AI KOL 主表",
            },
        ],
    }

    additional = [
        {
            "id": "papareact",
            "priority": "A",
            "name": "Papa React",
            "type": "Creator team and developer community",
            "contact": "https://www.youtube.com/@SonnySangha/about；https://www.papareact.com/",
            "officialUrl": "https://www.papareact.com/",
            "relationshipStatus": "未建联",
            "verification": "中",
            "nextAction": "确认是否代表 Sonny 处理商业合作，并询问是否有其他 AI/developer creators",
            "coverage": 1,
            "linkedKols": [kol_child(sonny, "kol:SonnySangha", "公开团队/社群入口，代理权限待核验")],
        },
        {
            "id": "sage-wave",
            "priority": "A",
            "name": "Sage Wave Media LLC",
            "type": "Creator-owned media company",
            "contact": "info@blogwithben.com",
            "officialUrl": "",
            "relationshipStatus": "未建联",
            "verification": "中",
            "nextAction": "确认是否只服务 Blog With Ben，或可提供其他 AI/productivity creator 资源",
            "coverage": 1,
            "linkedKols": [kol_child(blog_with_ben, "kol:blogwithben", "官网公开公司/邮箱，roster 范围待核验")],
        },
        {
            "id": "simon-schuster-speakers",
            "priority": "A",
            "name": "Simon & Schuster Speakers Bureau",
            "type": "Speaker representation",
            "contact": "erin.simpson@simonandschuster.com（Erin Simpson）",
            "officialUrl": "",
            "relationshipStatus": "未建联",
            "verification": "中高",
            "nextAction": "确认 Tiago Forte 的品牌内容和演讲合作范围，并询问 AI/tech speaker roster",
            "coverage": 1,
            "linkedKols": [kol_child(tiago, "kol:TiagoForte", "公开经纪联系人，品牌赞助授权范围待确认")],
        },
    ]

    return [
        lmg,
        smooth_media,
        electrify,
        futurepedia_network,
        kernel,
        claryo,
        passionfroot,
        markitup,
        *additional,
    ]


def set_title(ws, title: str, subtitle: str, end_col: int) -> None:
    end = get_column_letter(end_col)
    ws.merge_cells(f"A1:{end}2")
    ws["A1"] = title
    ws["A1"].font = Font(name="Aptos Display", size=22, bold=True, color=WHITE)
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].alignment = Alignment(vertical="center")
    ws.merge_cells(f"A3:{end}3")
    ws["A3"] = subtitle
    ws["A3"].font = Font(name="Aptos", size=10, color=MUTED)
    ws["A3"].fill = PatternFill("solid", fgColor=PALE_GRAY)
    ws["A3"].alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 28
    ws.row_dimensions[2].height = 12
    ws.row_dimensions[3].height = 34


def style_metric(ws, label_cell: str, value_cell: str, label: str, formula: str) -> None:
    ws[label_cell] = label
    ws[value_cell] = formula
    ws[label_cell].fill = PatternFill("solid", fgColor=NAVY)
    ws[label_cell].font = Font(name="Aptos", size=9, bold=True, color=WHITE)
    ws[label_cell].alignment = Alignment(horizontal="center", vertical="center")
    ws[value_cell].fill = PatternFill("solid", fgColor=PALE_TEAL)
    ws[value_cell].font = Font(name="Aptos Display", size=15, bold=True, color=NAVY)
    ws[value_cell].alignment = Alignment(horizontal="center", vertical="center")
    ws[value_cell].number_format = "0"


def style_header(cell) -> None:
    cell.fill = PatternFill("solid", fgColor=TEAL)
    cell.font = Font(name="Aptos", size=10, bold=True, color=WHITE)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = Border(bottom=Side(style="medium", color=NAVY))


def add_table(ws, ref: str, name: str) -> None:
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)


def set_external_link(cell, target: str) -> None:
    if not target:
        return
    cell.hyperlink = target
    cell.font = Font(name="Aptos", size=9, color=LINK, underline="single")


def set_internal_link(cell, location: str) -> None:
    cell._hyperlink = Hyperlink(ref=cell.coordinate, location=location, display=str(cell.value or ""))
    cell.font = Font(name="Aptos", size=9, color=LINK, underline="single")


def first_contact_link(value: str, fallback: str = "") -> str:
    email = re.search(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}", value or "")
    if email:
        return f"mailto:{email.group(0)}"
    url = re.search(r"https?://[^\s；;（）()]+", value or "")
    if url:
        return url.group(0)
    telegram = re.search(r"Telegram\s+@?([A-Za-z0-9_]+)", value or "", re.I)
    if telegram:
        return f"https://t.me/{telegram.group(1)}"
    whatsapp = re.search(r"WhatsApp\s+\+?([0-9 ]+)", value or "", re.I)
    if whatsapp:
        return f"https://wa.me/{whatsapp.group(1).replace(' ', '')}"
    return fallback


def style_data_block(ws, start_row: int, end_row: int, end_col: int, row_height: int) -> None:
    for row in range(start_row, end_row + 1):
        ws.row_dimensions[row].height = row_height
        for col in range(1, end_col + 1):
            cell = ws.cell(row, col)
            cell.font = Font(name="Aptos", size=9, color=DARK)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=THIN)


def build_kol_sheet(wb: Workbook, rows: list[dict[str, str]]) -> None:
    ws = wb.create_sheet("AI KOL 新建联")
    ws.sheet_properties.tabColor = CORAL
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 80
    ws.freeze_panes = "B8"
    set_title(
        ws,
        "AI KOL 新建联名单",
        "仅保留尚未联系且没有报价的候选。affiliate-only 账号已移除。S/A 优先级保留在备注中。编制日期 2026-09-06。",
        8,
    )

    first_data_row = 8
    last_data_row = first_data_row + len(rows) - 1
    style_metric(ws, "A4", "B4", "KOL 总数", f"=COUNTA(B{first_data_row}:B{last_data_row})")
    style_metric(ws, "C4", "D4", "S 级", f'=COUNTIF(H{first_data_row}:H{last_data_row},"S*")')
    style_metric(ws, "E4", "F4", "A 级", f'=COUNTIF(H{first_data_row}:H{last_data_row},"A*")')
    style_metric(ws, "G4", "H4", "未跟进", f'=COUNTIF(G{first_data_row}:G{last_data_row},"未跟进")')
    ws.row_dimensions[4].height = 27

    ws.merge_cells("A5:H5")
    ws["A5"] = (
        "本表用于首次建联。已报价、已有回复、affiliate-only 以及低质量集群不在本表。"
        "数据源：head_ai_kol_shortlist.csv、dashboard 底层 sponsorship/campaign/contact/quote 数据，以及用户指定后于 2026-09-06 核验的外部公开来源。"
    )
    ws["A5"].fill = PatternFill("solid", fgColor=PALE_YELLOW)
    ws["A5"].font = Font(name="Aptos", size=9, color=DARK)
    ws["A5"].alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[5].height = 32

    headers = ["平台", "账号", "主页链接", "粉丝数", "报价", "联系方式", "跟进时间", "备注"]
    for col, header in enumerate(headers, 1):
        ws.cell(7, col, header)
        style_header(ws.cell(7, col))
    ws.row_dimensions[7].height = 31

    for row_num, source in enumerate(rows, first_data_row):
        values = [source[h] for h in headers]
        for col, value in enumerate(values, 1):
            ws.cell(row_num, col, value)
        set_external_link(ws.cell(row_num, 3), source["主页链接"])
        contact_target = first_contact_link(source["联系方式"])
        if contact_target:
            set_external_link(ws.cell(row_num, 6), contact_target)

    style_data_block(ws, first_data_row, last_data_row, 8, 68)
    for row_num in range(first_data_row, last_data_row + 1):
        ws.cell(row_num, 4).alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)
        if "外部补充核验 2026-09-06" in str(ws.cell(row_num, 8).value or ""):
            ws.row_dimensions[row_num].height = 92
    add_table(ws, f"A7:H{last_data_row}", "NewAIKOLTable")

    full_range = f"A{first_data_row}:H{last_data_row}"
    ws.conditional_formatting.add(
        full_range,
        FormulaRule(formula=[f'LEFT($H{first_data_row},1)="S"'], fill=PatternFill("solid", fgColor=PALE_ORANGE)),
    )
    ws.conditional_formatting.add(
        f"E{first_data_row}:E{last_data_row}",
        FormulaRule(formula=[f'LEFT($E{first_data_row},3)="待询价"'], fill=PatternFill("solid", fgColor=PALE_YELLOW)),
    )
    ws.conditional_formatting.add(
        f"G{first_data_row}:G{last_data_row}",
        FormulaRule(formula=[f'$G{first_data_row}="未跟进"'], fill=PatternFill("solid", fgColor=PALE_YELLOW)),
    )

    widths = [12, 31, 43, 12, 48, 42, 28, 92]
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.auto_filter.ref = f"A7:H{last_data_row}"
    ws.print_title_rows = "1:7"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def supplier_permissions(name: str) -> str:
    mapping = {
        "Claryo Media": "AI creator sourcing、批量 campaign；creator 授权与交付能力待核验",
        "Passionfroot": "逐个 creator storefront 预订；通常不是独家代理，不等于 sub-affiliate 权限",
        "Linus Media Group": "集团自有媒体组合；适合大预算，不是开放 creator marketplace",
        "Papa React": "可能承接 creator 商务及开发者社群合作；是否有批量 roster 待确认",
        "Sage Wave Media LLC": "当前仅确认 Blog With Ben 的公司入口；其他 creator 资源待确认",
        "Simon & Schuster Speakers Bureau": "演讲与经纪预订；品牌内容和其他 AI/tech speaker 权限待确认",
        "Smooth Media": "creator 管理、品牌赞助、批量 creator 合作；sub-affiliate、白名单和跨平台授权需谈判",
        "Kernel Management / Reign Maker Talent": "科技 creator 经纪与品牌合作；完整 roster、批量采购和使用权待确认",
        "MarkitUp Media": "内容制作、YouTube 管理及 Ishan 商务入口；是否可代理其他 creator 待确认",
        "Electrify": "旗下 creator brands 的赞助与多创作者 campaign；并非开放 MCN/marketplace",
        "Futurepedia / Skill Leap AI": "三频道植入、专场、课程和 webinar；官网明确拒绝 affiliate offer",
    }
    return mapping.get(name, "合作权限待核验")


def build_supplier_sheet(wb: Workbook, data: dict) -> None:
    suppliers = data["supplierRows"]
    ws = wb.create_sheet("供应商新建联")
    ws.sheet_properties.tabColor = TEAL
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 80
    ws.freeze_panes = "B8"
    set_title(
        ws,
        "供应商新建联名单",
        "仅保留尚未联系或尚未取得报价的机构入口。点击“查看 KOL 明细”可跳转到本次候选 creator。编制日期 2026-09-06。",
        11,
    )

    summary_start = 8
    summary_end = summary_start + len(suppliers) - 1
    detail_header = summary_end + 4
    detail_start = detail_header + 1

    detail_rows: list[list[object]] = []
    first_detail_row: dict[str, int] = {}
    cursor = detail_start
    for supplier in suppliers:
        first_detail_row[supplier["id"]] = cursor
        for child in supplier["linkedKols"]:
            detail_rows.append([
                supplier["name"],
                child["creator"],
                child["platform"],
                child["handle"],
                child["profileUrl"],
                child["followers"],
                child["contact"],
                child["relationship"],
                child["creatorId"],
            ])
            cursor += 1
    detail_end = detail_start + len(detail_rows) - 1

    style_metric(ws, "A4", "B4", "供应商数", f"=COUNTA(B{summary_start}:B{summary_end})")
    style_metric(ws, "C4", "D4", "S 级入口", f'=COUNTIF(A{summary_start}:A{summary_end},"S")')
    style_metric(ws, "E4", "F4", "联系入口", f"=COUNTA(F{summary_start}:F{summary_end})")
    style_metric(ws, "G4", "H4", "低置信待核验", f'=COUNTIF(I{summary_start}:I{summary_end},"低")')
    ws["I4"] = "唯一 Creator 实体"
    ws["J4"] = f"=SUMPRODUCT(1/COUNTIF(I{detail_start}:I{detail_end},I{detail_start}:I{detail_end}))"
    ws["I4"].fill = PatternFill("solid", fgColor=NAVY)
    ws["I4"].font = Font(name="Aptos", size=9, bold=True, color=WHITE)
    ws["I4"].alignment = Alignment(horizontal="center", vertical="center")
    ws["J4"].fill = PatternFill("solid", fgColor=PALE_TEAL)
    ws["J4"].font = Font(name="Aptos Display", size=15, bold=True, color=NAVY)
    ws["J4"].alignment = Alignment(horizontal="center", vertical="center")
    ws["J4"].number_format = "0"
    ws["K4"] = "同一 creator 可能出现在多个公开入口中。此处按关联实体 ID 去重。"
    ws["K4"].font = Font(name="Aptos", size=8, italic=True, color=MUTED)
    ws["K4"].alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[4].height = 31

    ws.merge_cells("A5:K5")
    ws["A5"] = (
        "已有回复、已有报价、低置信集群和虚拟人入口均已移除。"
        "Passionfroot 仅列可预订候选样本，不作为独家供应商覆盖量；机构声称的总 roster 与本次已核验候选数分开记录。"
    )
    ws["A5"].fill = PatternFill("solid", fgColor=PALE_YELLOW)
    ws["A5"].font = Font(name="Aptos", size=9, color=DARK)
    ws["A5"].alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[5].height = 34

    headers = [
        "优先级", "供应商", "类型", "本次候选 KOL 数", "具体 KOL", "联系方式", "官方主页",
        "当前关系", "核验等级", "可谈合作权限", "下一步",
    ]
    for col, header in enumerate(headers, 1):
        ws.cell(7, col, header)
        style_header(ws.cell(7, col))
    ws.row_dimensions[7].height = 32

    supplier_row_map: dict[str, int] = {}
    for row_num, supplier in enumerate(suppliers, summary_start):
        supplier_row_map[supplier["id"]] = row_num
        values = [
            supplier["priority"], supplier["name"], supplier["type"], supplier["coverage"],
            f"查看 {supplier['coverage']} 个 KOL 明细 ↓", supplier["contact"], supplier["officialUrl"],
            supplier["relationshipStatus"], supplier["verification"], supplier_permissions(supplier["name"]),
            supplier["nextAction"],
        ]
        for col, value in enumerate(values, 1):
            ws.cell(row_num, col, value)
        detail_target = f"'供应商新建联'!A{first_detail_row[supplier['id']]}"
        set_internal_link(ws.cell(row_num, 5), detail_target)
        contact_target = first_contact_link(supplier["contact"], supplier["officialUrl"])
        if contact_target:
            set_external_link(ws.cell(row_num, 6), contact_target)
        if supplier["officialUrl"]:
            set_external_link(ws.cell(row_num, 7), supplier["officialUrl"])
    style_data_block(ws, summary_start, summary_end, 11, 58)
    for row_num in range(summary_start, summary_end + 1):
        ws.cell(row_num, 4).alignment = Alignment(horizontal="right", vertical="top")
    add_table(ws, f"A7:K{summary_end}", "SupplierSummaryTable")

    section_row = detail_header - 1
    ws.merge_cells(start_row=section_row, start_column=1, end_row=section_row, end_column=11)
    ws.cell(section_row, 1, "供应商关联 KOL 明细（从上表点击跳转；点击供应商名称返回上表）")
    ws.cell(section_row, 1).fill = PatternFill("solid", fgColor=NAVY)
    ws.cell(section_row, 1).font = Font(name="Aptos", size=11, bold=True, color=WHITE)
    ws.cell(section_row, 1).alignment = Alignment(vertical="center")
    ws.row_dimensions[section_row].height = 27

    detail_headers = ["供应商", "KOL", "平台", "账号", "主页链接", "粉丝数", "KOL 联系方式", "关系证据", "关联实体 ID"]
    for col, header in enumerate(detail_headers, 1):
        ws.cell(detail_header, col, header)
        style_header(ws.cell(detail_header, col))
    ws.row_dimensions[detail_header].height = 32

    supplier_name_to_id = {supplier["name"]: supplier["id"] for supplier in suppliers}
    for row_num, values in enumerate(detail_rows, detail_start):
        for col, value in enumerate(values, 1):
            ws.cell(row_num, col, value)
        supplier_id = supplier_name_to_id[values[0]]
        set_internal_link(ws.cell(row_num, 1), f"'供应商新建联'!A{supplier_row_map[supplier_id]}")
        if values[4]:
            set_external_link(ws.cell(row_num, 5), str(values[4]))
        contact_target = first_contact_link(str(values[6]))
        if contact_target:
            set_external_link(ws.cell(row_num, 7), contact_target)
        else:
            set_internal_link(ws.cell(row_num, 7), f"'供应商新建联'!F{supplier_row_map[supplier_id]}")
    style_data_block(ws, detail_start, detail_end, 9, 42)
    for row_num in range(detail_start, detail_end + 1):
        ws.cell(row_num, 6).alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)
    add_table(ws, f"A{detail_header}:I{detail_end}", "SupplierKOLDetailTable")

    ws.conditional_formatting.add(
        f"A{summary_start}:K{summary_end}",
        FormulaRule(formula=[f'$A{summary_start}="S"'], fill=PatternFill("solid", fgColor=PALE_ORANGE)),
    )
    ws.conditional_formatting.add(
        f"I{summary_start}:I{summary_end}",
        FormulaRule(formula=[f'$I{summary_start}="低"'], fill=PatternFill("solid", fgColor=PALE_RED)),
    )

    widths = [24, 28, 24, 15, 23, 38, 38, 22, 12, 58, 58]
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.print_title_rows = "1:7"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def validate(output: Path, expected_kols: int, expected_suppliers: int) -> dict[str, object]:
    wb = load_workbook(output, data_only=False)
    if wb.sheetnames != ["AI KOL 新建联", "供应商新建联"]:
        raise AssertionError(wb.sheetnames)
    kol = wb["AI KOL 新建联"]
    supplier = wb["供应商新建联"]
    expected_headers = ["平台", "账号", "主页链接", "粉丝数", "报价", "联系方式", "跟进时间", "备注"]
    actual_headers = [kol.cell(7, col).value for col in range(1, 9)]
    if actual_headers != expected_headers:
        raise AssertionError(actual_headers)
    if kol.max_row != 7 + expected_kols:
        raise AssertionError((kol.max_row, expected_kols))
    if supplier["B4"].value != f"=COUNTA(B8:B{7 + expected_suppliers})":
        raise AssertionError(supplier["B4"].value)

    blocked_suppliers = {
        "UnderCurrent", "Castle & Castle", "TalenTube", "PortilloBoss Network",
        "AI Planet Network", "Valid", "Neura Agency 标签集群",
    }
    for row in range(8, 8 + expected_suppliers):
        name = str(supplier.cell(row, 2).value or "")
        relationship = str(supplier.cell(row, 8).value or "")
        if name in blocked_suppliers or "已有" in relationship:
            raise AssertionError(f"Previously contacted supplier remained at row {row}: {name}")

    duplicate_profiles = []
    seen = set()
    missing = []
    invalid_links = []
    formula_errors = []
    for row in range(8, kol.max_row + 1):
        profile = str(kol.cell(row, 3).value or "").strip()
        if profile in seen:
            duplicate_profiles.append(profile)
        seen.add(profile)
        if not all(str(kol.cell(row, col).value or "").strip() for col in range(1, 9)):
            missing.append(row)
        if not re.match(r"^https?://", profile):
            invalid_links.append(profile)
        if not str(kol.cell(row, 5).value or "").startswith("待询价"):
            raise AssertionError(f"Quoted creator remained at row {row}")
        if str(kol.cell(row, 7).value or "") != "未跟进":
            raise AssertionError(f"Previously contacted creator remained at row {row}")
        if "affiliate-only" in str(kol.cell(row, 8).value or "").lower():
            raise AssertionError(f"Affiliate-only creator remained at row {row}")
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if isinstance(value, str) and value.strip().upper() in {
                    "#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!"
                }:
                    formula_errors.append(f"{ws.title}!{cell.coordinate}={value}")
    if duplicate_profiles or missing or invalid_links or formula_errors:
        raise AssertionError({
            "duplicate_profiles": duplicate_profiles,
            "missing_rows": missing,
            "invalid_links": invalid_links,
            "formula_errors": formula_errors,
        })
    return {
        "sheets": wb.sheetnames,
        "kol_rows": expected_kols,
        "supplier_rows": expected_suppliers,
        "supplier_detail_rows": supplier.max_row - (7 + expected_suppliers + 4),
        "duplicate_kol_profiles": len(duplicate_profiles),
        "missing_kol_cells": len(missing),
        "formula_errors": len(formula_errors),
        "tables": sum(len(ws.tables) for ws in wb.worksheets),
    }


def main() -> None:
    source_kols = load_kols()
    data = load_embedded_data()
    if len(source_kols) != data["stats"]["kolCount"]:
        raise AssertionError((len(source_kols), data["stats"]["kolCount"]))
    kols, removed_kols = filter_new_kols(source_kols, data)
    kols = merge_external_kols(kols)
    suppliers = build_new_supplier_rows(data, kols)
    filtered_data = {**data, "supplierRows": suppliers}

    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.creator = "Mango Labs"
    wb.properties.title = "AI KOL and supplier new outreach list"
    wb.properties.subject = "New AI creator and supplier outreach candidates without prior quotes or replies"
    wb.calculation.calcMode = "auto"
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True

    build_kol_sheet(wb, kols)
    build_supplier_sheet(wb, filtered_data)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    temp = OUTPUT_FILE.with_suffix(".tmp.xlsx")
    wb.save(temp)
    temp.replace(OUTPUT_FILE)
    result = validate(OUTPUT_FILE, len(kols), len(suppliers))
    print(json.dumps({
        "output": str(OUTPUT_FILE),
        "source_kols": len(source_kols),
        "removed_kols": len(removed_kols),
        **result,
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
