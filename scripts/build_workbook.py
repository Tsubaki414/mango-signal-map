#!/usr/bin/env python3
"""Build the operational Mango Labs AI BD pilot workbook from pipeline CSVs."""

from __future__ import annotations

import csv
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "outputs" / "pilot"
OUTPUT_DIR = ROOT / "outputs" / "mango_bd_pilot_20260824"
OUTPUT_FILE = OUTPUT_DIR / "Mango_Labs_AI_BD_Pilot.xlsx"

NAVY = "152238"
TEAL = "18A999"
CORAL = "FF6B5E"
PALE_TEAL = "E7F7F4"
PALE_BLUE = "EAF2FF"
PALE_YELLOW = "FFF4CC"
PALE_GREEN = "E8F5E9"
PALE_GRAY = "F4F6F8"
MID_GRAY = "D8DEE6"
DARK = "1F2937"
WHITE = "FFFFFF"
LINK = "0563C1"
THIN = Side(style="thin", color=MID_GRAY)


def read_csv(name: str) -> list[dict[str, str]]:
    with (SOURCE / name).open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def display(value: str | None) -> str:
    if value is None:
        return ""
    value = value.strip()
    if not value:
        return ""
    try:
        parsed = json.loads(value)
    except (json.JSONDecodeError, TypeError):
        return value
    if isinstance(parsed, list):
        return " • ".join(str(item) for item in parsed)
    if isinstance(parsed, dict):
        return " • ".join(f"{key}: {val}" for key, val in parsed.items())
    return str(parsed)


def as_number(value: str, default: float = 0.0) -> float:
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def title_band(ws, title: str, subtitle: str, end_col: int) -> None:
    end = ws.cell(row=1, column=end_col).column_letter
    ws.merge_cells(f"A1:{end}2")
    cell = ws["A1"]
    cell.value = title
    cell.font = Font(name="Aptos Display", size=22, bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(vertical="center")
    ws.merge_cells(f"A3:{end}3")
    ws["A3"] = subtitle
    ws["A3"].font = Font(name="Aptos", size=10, color="52606D")
    ws["A3"].fill = PatternFill("solid", fgColor=PALE_GRAY)
    ws["A3"].alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 28
    ws.row_dimensions[2].height = 12
    ws.row_dimensions[3].height = 30


def style_sheet(ws, freeze: str = "A6", zoom: int = 90) -> None:
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = freeze
    ws.sheet_view.zoomScale = zoom
    ws.auto_filter.ref = None


def write_table(
    ws,
    headers: list[str],
    rows: list[list[object]],
    widths: list[int],
    table_name: str,
    start_row: int = 5,
    row_height: int = 48,
) -> None:
    for col, header in enumerate(headers, 1):
        cell = ws.cell(start_row, col, header)
        cell.font = Font(name="Aptos", size=10, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=TEAL)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="medium", color=NAVY))
    ws.row_dimensions[start_row].height = 32

    for r_offset, row in enumerate(rows, start_row + 1):
        for col, value in enumerate(row, 1):
            cell = ws.cell(r_offset, col, value)
            cell.font = Font(name="Aptos", size=9, color=DARK)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=THIN)
            if r_offset % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=PALE_GRAY)
        ws.row_dimensions[r_offset].height = row_height

    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width

    if rows:
        end_row = start_row + len(rows)
        end_col = ws.cell(start_row, len(headers)).column_letter
        table = Table(displayName=table_name, ref=f"A{start_row}:{end_col}{end_row}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        ws.add_table(table)


def add_link(cell, url: str) -> None:
    if url:
        cell.hyperlink = url
        cell.font = Font(name="Aptos", size=9, color=LINK, underline="single")


def build_dashboard(wb: Workbook, projects: list[dict], actions: list[dict]) -> None:
    ws = wb.create_sheet("Dashboard")
    ws.sheet_properties.tabColor = CORAL
    title_band(
        ws,
        "Mango Labs · AI BD Intelligence Pilot",
        "Operating dashboard · 2026-08-24 snapshot · formulas recalculate when the workbook opens · confidence and relationship strength are kept separate.",
        10,
    )
    style_sheet(ws, freeze="A12", zoom=95)
    for col in range(1, 11):
        ws.column_dimensions[get_column_letter(col)].width = 16

    cards = [
        ("A5:B5", "A6:B7", "TARGET COMPANIES", "=COUNTA(Projects!$B$6:$B$17)"),
        ("C5:D5", "C6:D7", "RELATIONSHIP EDGES", "=COUNTA(Relationships!$A$6:$A$205)"),
        ("E5:F5", "E6:F7", "EVIDENCE ITEMS", "=COUNTA(Evidence!$A$6:$A$205)"),
        ("G5:H5", "G6:H7", "SPONSORSHIP RECORDS", "=COUNTA(Creators!$A$6:$A$105)"),
        ("I5:J5", "I6:J7", "VERIFIED 2-HOP TARGETS", '=COUNTIF(Projects!$G$6:$G$17,"verified_two_hop")'),
    ]
    for label_range, value_range, label, formula in cards:
        ws.merge_cells(label_range)
        ws.merge_cells(value_range)
        label_cell = ws[label_range.split(":")[0]]
        value_cell = ws[value_range.split(":")[0]]
        label_cell.value = label
        value_cell.value = formula
        label_cell.fill = PatternFill("solid", fgColor=NAVY)
        label_cell.font = Font(name="Aptos", size=9, bold=True, color=WHITE)
        label_cell.alignment = Alignment(horizontal="center", vertical="center")
        value_cell.fill = PatternFill("solid", fgColor=PALE_TEAL)
        value_cell.font = Font(name="Aptos Display", size=24, bold=True, color=NAVY)
        value_cell.alignment = Alignment(horizontal="center", vertical="center")
        value_cell.number_format = "0"

    ws.merge_cells("A9:J9")
    ws["A9"] = "PRIORITY PIPELINE — SCORING IS FORMULA-DRIVEN FROM THE PROJECTS SHEET"
    ws["A9"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A9"].font = Font(name="Aptos", size=10, bold=True, color=WHITE)
    ws["A9"].alignment = Alignment(vertical="center")

    dashboard_headers = ["Rank", "Company", "Score", "Reachability", "Immediate move"]
    for col, header in enumerate(dashboard_headers, 1):
        cell = ws.cell(11, col, header)
        cell.fill = PatternFill("solid", fgColor=TEAL)
        cell.font = Font(name="Aptos", size=10, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for idx in range(min(8, len(projects))):
        row = 12 + idx
        project_row = 6 + idx
        action = actions[idx] if idx < len(actions) else {}
        ws.cell(row, 1, f"=Projects!A{project_row}")
        ws.cell(row, 2, f"=Projects!B{project_row}")
        ws.cell(row, 3, f"=Projects!N{project_row}")
        ws.cell(row, 4, f"=Projects!G{project_row}")
        ws.cell(row, 5, action.get("first_outreach_angle", ""))
        for col in range(1, 6):
            cell = ws.cell(row, col)
            cell.font = Font(name="Aptos", size=9, color=DARK)
            cell.fill = PatternFill("solid", fgColor=PALE_GRAY if row % 2 == 0 else WHITE)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=THIN)
        ws.row_dimensions[row].height = 46
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 10
    ws.column_dimensions["D"].width = 23
    ws.column_dimensions["E"].width = 74

    chart = BarChart()
    chart.type = "bar"
    chart.style = 10
    chart.title = "Top 8 weighted priorities"
    chart.y_axis.title = "Company"
    chart.x_axis.title = "Score / 100"
    chart.height = 7.7
    chart.width = 13.8
    chart.legend = None
    data = Reference(ws, min_col=3, min_row=11, max_row=19)
    cats = Reference(ws, min_col=2, min_row=12, max_row=19)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.x_axis.scaling.max = 100
    ws.add_chart(chart, "G11")

    ws.merge_cells("A22:J22")
    ws["A22"] = "RELATIONSHIP READOUT"
    ws["A22"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A22"].font = Font(name="Aptos", size=10, bold=True, color=WHITE)
    notes = [
        "Warmest verified operating paths: Mango → Jennie Liu → Chang Chen → Gamma; Mango → Jennie Liu → Alec Wilcock → ElevenLabs; Mango → Jennie Liu → Jeddi Mees → HeyGen.",
        "Weak X-only signals: Jennie follows Replit, Cursor, Perplexity, Pika and Runway. A follow is not evidence of a working relationship or willingness to introduce.",
        "No direct target-company follow was found across the fully paginated X following IDs of Amy, Doris, Dov or Solomon among the seven resolved target accounts.",
    ]
    for offset, note in enumerate(notes, 23):
        ws.merge_cells(start_row=offset, start_column=1, end_row=offset, end_column=10)
        ws.cell(offset, 1, f"• {note}")
        ws.cell(offset, 1).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(offset, 1).font = Font(name="Aptos", size=10, color=DARK)
        ws.cell(offset, 1).fill = PatternFill("solid", fgColor=PALE_YELLOW if offset == 24 else PALE_GRAY)
        ws.row_dimensions[offset].height = 38


def build_scoring(wb: Workbook) -> None:
    ws = wb.create_sheet("Scoring")
    ws.sheet_properties.tabColor = NAVY
    title_band(ws, "Scoring & confidence model", "Weights are editable inputs. Projects!Overall Priority recalculates from these cells.", 5)
    style_sheet(ws, freeze="A5")
    headers = ["Metric", "Weight", "Why it matters", "Input owner", "Audit note"]
    rows = [
        ["Reachability", 0.35, "Largest weight: access determines whether a strong target can become a live conversation.", "BD lead", "Scores relationship evidence, not prestige."],
        ["Budget", 0.25, "Uses explicit ARR/funding, active programs, job briefs and disclosed paid activity.", "Research", "Confirmed > high probability > speculative."],
        ["Fit", 0.20, "Mango's regional creator, localization, campaign and relationship advantage.", "Growth", "0–100 analyst input."],
        ["Timing", 0.10, "Recent raise, launch, expansion, hiring, or active campaign window.", "Research", "0–100 analyst input."],
        ["Evidence quality", 0.10, "Source quality and triangulation depth.", "Data", "Official + disclosed campaign evidence scores highest."],
    ]
    write_table(ws, headers, rows, [24, 13, 72, 18, 52], "ScoringWeights", row_height=42)
    for row in range(6, 11):
        ws.cell(row, 2).number_format = "0%"
        ws.cell(row, 2).fill = PatternFill("solid", fgColor=PALE_BLUE)
    ws["B6"].comment = Comment("Editable model input. Reachability deliberately has the largest single weight.", "User")

    ws["A13"] = "Reachability level"
    ws["B13"] = "Score"
    ws["C13"] = "Interpretation"
    for cell in ws[13][:3]:
        cell.fill = PatternFill("solid", fgColor=TEAL)
        cell.font = Font(name="Aptos", size=10, bold=True, color=WHITE)
    reach_map = [
        ["verified_direct", 100, "Verified direct operating relationship"],
        ["verified_two_hop", 68, "Verified connector-to-operator path; intro willingness still requires validation"],
        ["public_business_contact", 45, "Relevant published program or business contact"],
        ["cold_only", 10, "No verified operator path"],
    ]
    for row_idx, row in enumerate(reach_map, 14):
        for col_idx, value in enumerate(row, 1):
            ws.cell(row_idx, col_idx, value)
            ws.cell(row_idx, col_idx).border = Border(bottom=THIN)
            ws.cell(row_idx, col_idx).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row_idx].height = 32


def build_projects(wb: Workbook, projects: list[dict]) -> None:
    ws = wb.create_sheet("Projects")
    ws.sheet_properties.tabColor = CORAL
    title_band(ws, "Target project database", "Blue cells are analyst inputs; green cells are formulas; evidence IDs and paths keep every score auditable.", 23)
    style_sheet(ws, freeze="B6", zoom=80)
    headers = [
        "Rank", "Company", "Category", "Geography", "Stage", "Business Model", "Reachability Level",
        "Reachability Score", "Budget Confidence", "Budget Score", "Fit Score", "Timing Score",
        "Evidence Quality", "Overall Priority", "Mango Relationship Path", "Decision Makers", "Public Contact Methods",
        "Budget Signals", "Marketing Channels", "Historical Campaigns", "Recommended Next Action", "Evidence IDs", "Last Verified",
    ]
    widths = [8, 18, 27, 29, 20, 31, 24, 16, 19, 14, 12, 12, 16, 15, 37, 35, 38, 48, 36, 48, 68, 45, 15]
    rows = []
    for project in projects:
        rows.append([
            "", project["company"], project["category"], project["geography"], project["stage"], project["business_model"],
            project["reachability_level"], as_number(project["reachability_score"]), project["budget_confidence"],
            as_number(project["budget_score"]), as_number(project["fit_score"]), as_number(project["timing_score"]),
            as_number(project["evidence_quality_score"]), "", display(project["mango_relationship_path"]),
            display(project["decision_maker_ids"]), display(project["public_contact_methods"]), display(project["budget_signals"]),
            display(project["marketing_channels"]), display(project["historical_campaigns"]), project["recommended_next_action"],
            display(project["evidence_ids"]), project["last_verified_date"],
        ])
    write_table(ws, headers, rows, widths, "TargetProjects", row_height=76)
    end_row = 5 + len(rows)
    for row in range(6, end_row + 1):
        ws.cell(row, 1, f"=RANK.EQ(N{row},$N$6:$N${end_row},0)")
        ws.cell(row, 14, f"=ROUND(H{row}*Scoring!$B$6+J{row}*Scoring!$B$7+K{row}*Scoring!$B$8+L{row}*Scoring!$B$9+M{row}*Scoring!$B$10,1)")
        for col in (8, 10, 11, 12, 13):
            ws.cell(row, col).fill = PatternFill("solid", fgColor=PALE_BLUE)
            ws.cell(row, col).font = Font(name="Aptos", size=9, color=LINK)
        for col in (1, 14):
            ws.cell(row, col).fill = PatternFill("solid", fgColor=PALE_GREEN)
            ws.cell(row, col).font = Font(name="Aptos", size=9, bold=True, color="18733C")
        ws.cell(row, 14).number_format = "0.0"
    ws["N5"].comment = Comment("Formula: Reachability 35% + Budget 25% + Fit 20% + Timing 10% + Evidence 10%.", "User")
    ws["O5"].comment = Comment("A path records observed reachability only. It does not imply the connector has agreed to introduce Mango.", "User")
    ws.conditional_formatting.add(f"N6:N{end_row}", ColorScaleRule(start_type="num", start_value=0, start_color="F8696B", mid_type="num", mid_value=70, mid_color="FFEB84", end_type="num", end_value=100, end_color="63BE7B"))


def build_actions(wb: Workbook, actions: list[dict]) -> None:
    ws = wb.create_sheet("Priority Actions")
    ws.sheet_properties.tabColor = CORAL
    title_band(ws, "Priority action queue", "The first eight moves are ordered for execution. Yellow columns are operating fields for Mango to update.", 14)
    style_sheet(ws, freeze="D6", zoom=80)
    headers = ["Rank", "Project", "Status", "Next Follow-Up", "Mango Owner", "Approach Mode", "Contact", "Intro Path", "Why Now", "Budget Case", "Mango Value", "First Outreach Angle", "Fallback Path", "Notes"]
    widths = [8, 18, 16, 17, 19, 29, 37, 39, 53, 53, 53, 70, 63, 45]
    rows = []
    for action in actions:
        rows.append([
            int(action["rank"]), action["project"], "Not started", "", action["mango_owner"], action["approach_mode"],
            action["contact"], action["intro_path"], action["why_now"], action["budget_case"], action["mango_value"],
            action["first_outreach_angle"], action["fallback_path"], "",
        ])
    write_table(ws, headers, rows, widths, "PriorityActionQueue", row_height=86)
    end_row = 5 + len(rows)
    status = DataValidation(type="list", formula1='"Not started,Intro requested,Contacted,Meeting booked,Proposal sent,Won,Paused,Closed"', allow_blank=False)
    ws.add_data_validation(status)
    status.add(f"C6:C{end_row}")
    for row in range(6, end_row + 1):
        for col in (3, 4, 14):
            ws.cell(row, col).fill = PatternFill("solid", fgColor=PALE_YELLOW)
        ws.cell(row, 4).number_format = "yyyy-mm-dd"
    ws["C5"].comment = Comment("Editable pipeline stage. Use the dropdown rather than free text.", "User")


def build_generic_sheet(wb: Workbook, name: str, title: str, subtitle: str, headers: list[str], rows: list[list[object]], widths: list[int], table_name: str, row_height: int = 54) -> None:
    ws = wb.create_sheet(name)
    ws.sheet_properties.tabColor = TEAL
    title_band(ws, title, subtitle, len(headers))
    style_sheet(ws, freeze="A6", zoom=80)
    write_table(ws, headers, rows, widths, table_name, row_height=row_height)


def build_readme(wb: Workbook) -> None:
    ws = wb.create_sheet("README")
    ws.sheet_properties.tabColor = NAVY
    title_band(ws, "Workbook guide & data dictionary", "This workbook is an evidence-backed BD operating artifact, not a claim that any connector has agreed to help.", 6)
    style_sheet(ws, freeze="A5", zoom=95)
    headers = ["Section", "Field / Sheet", "Definition", "How to use", "Confidence rule", "Refresh cadence"]
    rows = [
        ["Workflow", "Priority Actions", "Execution queue for the eight highest-value near-term moves.", "Assign owner, status and follow-up; log outreach in Notes.", "Action language never upgrades weak relationship evidence.", "Weekly"],
        ["Model", "Projects", "12 target companies with GTM, operator, budget and relationship fields.", "Filter by reachability, score, category or timing.", "Overall score is formula-driven from Scoring weights.", "Monthly / event-driven"],
        ["Relationship", "Relationships", "Directional and non-directional graph edges with strength and evidence IDs.", "Trace Mango → connector → operator/company before asking for an intro.", "X follow = weak; LinkedIn connection = verified path, not consent.", "Monthly"],
        ["Campaign", "Creators", "Disclosed sponsorships, repeat buys, affiliate or event partnerships.", "Use creators with real campaign history as proof and possible connectors.", "Confirmed requires explicit disclosure or official program evidence.", "Monthly"],
        ["Source", "Evidence", "Atomic claims with URL, date, source type and excerpt.", "Open the URL before externalizing a claim or sending a proposal.", "Official > disclosed campaign > reputable secondary > directional signal.", "Quarterly / before outreach"],
        ["X data", "Rapid X", "All X profile/following data in this pilot was collected via Rapid X API.", "Set RAPID_X_API_KEY in local .env; never paste it into source or workbook.", "Full pagination used for seed-account following IDs.", "On rerun"],
        ["Limitation", "Unresolved X handles", "Five candidate target handles did not resolve in Rapid X during collection.", "Treat absence as unknown, not as proof of no account/follow.", "No inference from failed resolution.", "On rerun"],
        ["Security", "API credential", "The supplied screenshot visibly exposed a RapidAPI credential.", "Rotate the key after this run and keep .env gitignored.", "Credential is never included in outputs.", "Immediate"],
    ]
    write_table(ws, headers, rows, [19, 27, 62, 64, 57, 22], "WorkbookReadme", row_height=64)


def main() -> None:
    projects = sorted(read_csv("projects.csv"), key=lambda x: as_number(x["overall_priority"]), reverse=True)
    actions = sorted(read_csv("priority_actions.csv"), key=lambda x: int(x["rank"]))
    creators = read_csv("creators_sponsorships.csv")
    relationships = read_csv("relationships.csv")
    entities = read_csv("entities.csv")
    evidence = read_csv("evidence.csv")
    cases = read_csv("gtm_case_studies.csv")

    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.creator = "Mango Labs"
    wb.properties.lastModifiedBy = "User"
    wb.properties.title = "Mango Labs AI BD Intelligence Pilot"
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"

    build_dashboard(wb, projects, actions)
    build_scoring(wb)
    build_projects(wb, projects)
    build_actions(wb, actions)

    creator_headers = ["ID", "Creator ID", "Project ID", "Platform", "Content Title", "Content URL", "Content Date", "Cooperation Type", "Disclosure", "Repeated", "Audience Fit", "Relationship Value", "Contact Method", "Evidence IDs", "Confidence"]
    creator_rows = [[r["id"], r["creator_id"], r["project_id"], r["platform"], r["content_title"], r["content_url"], r["content_date"], r["cooperation_type"], r["disclosure"], r["repeated"], r["audience_fit"], r["relationship_value"], r["contact_method"], display(r["evidence_ids"]), r["confidence"]] for r in creators]
    build_generic_sheet(wb, "Creators", "Creator & sponsorship evidence", "Confirmed sponsorships are separated from affiliate-only or high-probability signals.", creator_headers, creator_rows, [24, 25, 20, 13, 54, 48, 15, 29, 42, 12, 40, 48, 37, 35, 16], "CreatorSponsorships", 68)
    for row, r in enumerate(creators, 6):
        add_link(wb["Creators"].cell(row, 6), r["content_url"])

    rel_headers = ["ID", "Source ID", "Target ID", "Relationship Type", "Strength", "Confidence", "Evidence IDs", "Observed Date", "Directional"]
    rel_rows = [[r["id"], r["source_id"], r["target_id"], r["relationship_type"], as_number(r["strength"]), r["confidence"], display(r["evidence_ids"]), r["observed_date"], r["directional"]] for r in relationships]
    build_generic_sheet(wb, "Relationships", "Relationship graph edges", "Strength is a routing heuristic. It must not be read as consent, endorsement or intro willingness.", rel_headers, rel_rows, [30, 25, 26, 33, 12, 16, 35, 16, 13], "RelationshipEdges", 34)
    wb["Relationships"].conditional_formatting.add(f"E6:E{5+len(rel_rows)}", ColorScaleRule(start_type="num", start_value=0, start_color="F8696B", mid_type="num", mid_value=50, mid_color="FFEB84", end_type="num", end_value=100, end_color="63BE7B"))

    entity_headers = ["ID", "Kind", "Name", "Handles", "Attributes"]
    entity_rows = [[r["id"], r["kind"], r["name"], display(r["handles"]), display(r["attributes"])] for r in entities]
    build_generic_sheet(wb, "Entities", "Entity registry", "Normalized registry for Mango, people, target companies, operators and creators.", entity_headers, entity_rows, [30, 16, 30, 48, 86], "EntityRegistry", 42)

    evidence_headers = ["ID", "Claim", "Source URL", "Source Title", "Published Date", "Verified Date", "Confidence", "Source Type", "Evidence Excerpt"]
    evidence_rows = [[r["id"], r["claim"], r["source_url"], r["source_title"], r["published_date"], r["verified_date"], r["confidence"], r["source_type"], r["excerpt"]] for r in evidence]
    build_generic_sheet(wb, "Evidence", "Evidence ledger", "Each material claim is traceable to a dated source. Open and re-check before using claims externally.", evidence_headers, evidence_rows, [31, 72, 54, 49, 16, 16, 16, 18, 72], "EvidenceLedger", 76)
    for row, r in enumerate(evidence, 6):
        add_link(wb["Evidence"].cell(row, 3), r["source_url"])

    case_headers = ["Company", "Channel Mix", "Creator Tiers", "Campaign Rhythm", "Partner Program", "Operator Model", "Selection Logic", "Relationship Compounding", "Mango Takeaway"]
    case_rows = [[r["company"], r["channel_mix"], r["creator_tiers"], r["campaign_rhythm"], r["partner_program"], r["operator_model"], r["selection_logic"], r["relationship_compounding"], r["mango_takeaway"]] for r in cases]
    build_generic_sheet(wb, "GTM Cases", "Reverse-engineered GTM cases", "Mature programs reveal operating systems—not just channel lists. Use these patterns to shape Mango's wedge.", case_headers, case_rows, [17, 52, 45, 44, 46, 47, 53, 54, 58], "GTMCases", 100)

    build_readme(wb)
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT_FILE)
    print(OUTPUT_FILE)


if __name__ == "__main__":
    main()
